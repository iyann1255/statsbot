import asyncio, os
from datetime import datetime, timezone, timedelta
from io import BytesIO

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from aiogram import Bot, Dispatcher, F
from aiogram.types import Message, BufferedInputFile, InlineKeyboardMarkup, InlineKeyboardButton
from aiogram.filters import Command
from aiogram.enums import ParseMode

from config import BOT_TOKEN
from database import (init_db, record_message, get_top_members,
                      get_daily_totals, get_hourly_stats,
                      get_user_stats, get_chat_total, get_monthly_stats,
                      get_unique_users, get_users_monthly_total, get_db_path)

bot = Bot(token=BOT_TOKEN)
dp  = Dispatcher()

WIB = timezone(timedelta(hours=7))

def now_wib():
    return datetime.now(WIB)

def esc(t: str) -> str:
    for ch in r'\_*[]()~`>#+-=|{}.!':
        t = t.replace(ch, f'\\{ch}')
    return t

# ── /top ─────────────────────────────────────────────────────────────
@dp.message(Command("top"))
async def cmd_top(msg: Message):
    if msg.chat.type not in ("group", "supergroup"):
        return await msg.reply("Perintah ini hanya untuk grup\\.", parse_mode="MarkdownV2")
    if not msg.from_user:
        return

    member = await bot.get_chat_member(msg.chat.id, msg.from_user.id)
    limit = 50 if member.status in ("administrator", "creator") else 10

    rows = get_top_members(msg.chat.id, limit=limit, days=30)
    if not rows:
        return await msg.reply("Belum ada data pesan\\.", parse_mode="MarkdownV2")
    medals = ["🥇","🥈","🥉"] + [f"{i}\\." for i in range(4, limit + 1)]
    lines = []
    for i, r in enumerate(rows):
        name = esc(r["full_name"] or r["username"] or str(r["user_id"]))
        lines.append(f"{medals[i]} {name} — *{esc(str(r['total']))}* pesan")
    await msg.reply(
        f"🏆 *Top {limit} Member \\(30 hari terakhir\\)*\n\n" + "\n".join(lines),
        parse_mode="MarkdownV2"
    )

# ── /stat ─────────────────────────────────────────────────────────────
@dp.message(Command("stat"))
async def cmd_stat(msg: Message):
    if msg.chat.type not in ("group", "supergroup"):
        return await msg.reply("Perintah ini hanya untuk grup\\.", parse_mode="MarkdownV2")
    if not msg.from_user:
        return
    target = msg.reply_to_message.from_user if msg.reply_to_message else msg.from_user
    if target.is_bot:
        return await msg.reply("Bot tidak punya statistik\\.", parse_mode="MarkdownV2")
    row, rank = get_user_stats(msg.chat.id, target.id)
    if not row or row["total"] is None:
        name = esc(target.full_name or target.username or str(target.id))
        return await msg.reply(f"👤 *{name}* belum punya data pesan\\.", parse_mode="MarkdownV2")
    name = esc(target.full_name or target.username or str(target.id))
    uname = f"@{esc(target.username)}" if target.username else "\\-"
    total = row["total"] or 0
    week  = row["week"]  or 0
    today = row["today"] or 0
    await msg.reply(
        f"👤 *{name}*\n"
        f"{uname}\n\n"
        f"Total pesan: *{esc(str(total))}*\n"
        f"7 hari ini: *{esc(str(week))}*\n"
        f"Hari ini: *{esc(str(today))}*\n"
        f"Peringkat: *\\#{esc(str(rank))}*",
        parse_mode="MarkdownV2"
    )

# ── /statadmin ────────────────────────────────────────────────────────
@dp.message(Command("statadmin"))
async def cmd_statadmin(msg: Message):
    if msg.chat.type not in ("group", "supergroup"):
        return await msg.reply("Perintah ini hanya untuk grup\\.", parse_mode="MarkdownV2")
    if not msg.from_user:
        return
    member = await bot.get_chat_member(msg.chat.id, msg.from_user.id)
    if member.status not in ("administrator", "creator"):
        return await msg.reply("Hanya admin yang bisa pakai perintah ini\\.", parse_mode="MarkdownV2")

    now = now_wib()
    admins = await bot.get_chat_administrators(msg.chat.id)
    admin_ids = [a.user.id for a in admins if not a.user.is_bot]
    monthly = get_users_monthly_total(msg.chat.id, now.year, now.month, admin_ids)

    lines = []
    for i, a in enumerate(admins, 1):
        if a.user.is_bot:
            continue
        aname = esc(a.user.full_name or a.user.username or str(a.user.id))
        pesan = monthly.get(a.user.id, 0)
        lines.append(f"{i}\\. {aname} — *{esc(str(pesan))}* pesan")

    month_esc = esc(now.strftime("%B %Y"))
    await msg.reply(
        f"📋 *Statistik Admin — {month_esc}*\n\n" + "\n".join(lines),
        parse_mode="MarkdownV2"
    )

# ── /grupstat ─────────────────────────────────────────────────────────
@dp.message(Command("grupstat"))
async def cmd_grupstat(msg: Message):
    if msg.chat.type not in ("group", "supergroup"):
        return await msg.reply("Perintah ini hanya untuk grup\\.", parse_mode="MarkdownV2")
    total = get_chat_total(msg.chat.id)
    rows7 = get_daily_totals(msg.chat.id, days=7)
    week  = sum(r["total"] for r in rows7)
    today_str = now_wib().strftime("%Y-%m-%d")
    today = next((r["total"] for r in rows7 if r["date"] == today_str), 0)
    await msg.reply(
        f"📊 *Statistik Grup*\n\n"
        f"Total pesan: *{esc(str(total))}*\n"
        f"7 hari ini: *{esc(str(week))}*\n"
        f"Hari ini: *{esc(str(today))}*",
        parse_mode="MarkdownV2"
    )

# ── /grafik ───────────────────────────────────────────────────────────
@dp.message(Command("grafik"))
async def cmd_grafik(msg: Message):
    if msg.chat.type not in ("group", "supergroup"):
        return await msg.reply("Perintah ini hanya untuk grup\\.", parse_mode="MarkdownV2")
    rows = get_daily_totals(msg.chat.id, days=7)
    if not rows:
        return await msg.reply("Belum ada data pesan\\.", parse_mode="MarkdownV2")

    dates  = [r["date"][5:] for r in rows]   # MM-DD
    totals = [r["total"] for r in rows]

    fig, ax = plt.subplots(figsize=(8, 4))
    ax.bar(dates, totals, color="#5865F2", width=0.5)
    ax.set_title("Aktivitas Pesan 7 Hari Terakhir", fontsize=13, pad=10)
    ax.set_xlabel("Tanggal"); ax.set_ylabel("Jumlah Pesan")
    ax.yaxis.set_major_locator(ticker.MaxNLocator(integer=True))
    ax.grid(axis="y", alpha=0.3)
    plt.tight_layout()

    buf = BytesIO()
    plt.savefig(buf, format="png", dpi=120)
    plt.close(fig)
    buf.seek(0)
    await msg.reply_photo(BufferedInputFile(buf.read(), "grafik.png"),
                          caption="📈 Aktivitas pesan 7 hari terakhir")

# ── /jam ──────────────────────────────────────────────────────────────
@dp.message(Command("jam"))
async def cmd_jam(msg: Message):
    if msg.chat.type not in ("group", "supergroup"):
        return await msg.reply("Perintah ini hanya untuk grup\\.", parse_mode="MarkdownV2")
    rows = get_hourly_stats(msg.chat.id, days=7)
    if not rows:
        return await msg.reply("Belum ada data pesan\\.", parse_mode="MarkdownV2")

    hours  = [r["hour"] for r in rows]
    totals = [r["total"] for r in rows]
    peak_h = hours[totals.index(max(totals))]

    fig, ax = plt.subplots(figsize=(10, 4))
    ax.bar([f"{h:02d}:00" for h in hours], totals, color="#57F287", width=0.6)
    ax.set_title("Aktivitas per Jam (7 Hari Terakhir, WIB)", fontsize=13, pad=10)
    ax.set_xlabel("Jam"); ax.set_ylabel("Jumlah Pesan")
    ax.yaxis.set_major_locator(ticker.MaxNLocator(integer=True))
    ax.grid(axis="y", alpha=0.3)
    plt.xticks(rotation=45, fontsize=8)
    plt.tight_layout()

    buf = BytesIO()
    plt.savefig(buf, format="png", dpi=120)
    plt.close(fig)
    buf.seek(0)
    await msg.reply_photo(BufferedInputFile(buf.read(), "jam.png"),
                          caption=f"🕐 Jam tersibuk: *{peak_h:02d}:00 WIB*",
                          parse_mode="MarkdownV2")

# ── Keyboards ────────────────────────────────────────────────────────
def kb_main(username: str) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="Tambah ke Grup", url=f"https://t.me/{username}?startgroup=true")],
        [InlineKeyboardButton(text="Daftar Perintah", callback_data="show_commands"),
         InlineKeyboardButton(text="Owner", url="tg://user?id=568033927")],
        [InlineKeyboardButton(text="Info Update", url="https://t.me/oneonlysepp")],
    ])

def kb_back() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="Kembali", callback_data="back_start")
    ]])

TEKS_START = (
    "📊 *StatBot*\n\n"
    "Pantau aktivitas chat grup kamu secara otomatis\\.\n\n"
    "• Tidak perlu jadi admin\n"
    "• Rekam pesan semua member\n"
    "• Grafik harian & per jam\n"
    "• Ranking, statistik personal & grup"
)

TEKS_COMMANDS = (
    "📋 *Daftar Perintah*\n\n"
    "/top — Ranking member teraktif \\(30 hari\\)\n"
    "/stat — Statistik pesanmu \\(reply = cek orang lain\\)\n"
    "/grupstat — Ringkasan aktivitas grup\n"
    "/grafik — Grafik 7 hari terakhir\n"
    "/jam — Aktivitas per jam \\(WIB\\)\n"
    "/statadmin — Statistik admin bulan ini\n"
    "/export — Export data ke Excel\n\n"
    "Tambahkan ke grup, langsung jalan\\."
)

# ── /start ───────────────────────────────────────────────────────────
@dp.message(Command("start"))
async def cmd_start(msg: Message):
    me = await bot.get_me()
    await msg.answer(TEKS_START, parse_mode="MarkdownV2", reply_markup=kb_main(me.username))

# ── /help ─────────────────────────────────────────────────────────────
@dp.message(Command("help"))
async def cmd_help(msg: Message):
    await msg.answer(TEKS_COMMANDS, parse_mode="MarkdownV2", reply_markup=kb_back())

from aiogram.types import CallbackQuery

@dp.callback_query(F.data == "show_commands")
async def cb_show_commands(cb: CallbackQuery):
    await cb.message.edit_text(TEKS_COMMANDS, parse_mode="MarkdownV2", reply_markup=kb_back())
    await cb.answer()

@dp.callback_query(F.data == "back_start")
async def cb_back_start(cb: CallbackQuery):
    me = await bot.get_me()
    await cb.message.edit_text(TEKS_START, parse_mode="MarkdownV2", reply_markup=kb_main(me.username))
    await cb.answer()

# ── /export ───────────────────────────────────────────────────────────
@dp.message(Command("export"))
async def cmd_export(msg: Message):
    if msg.chat.type not in ("group", "supergroup"):
        return await msg.reply("Perintah ini hanya untuk grup\\.", parse_mode="MarkdownV2")
    if not msg.from_user:
        return

    # Cek apakah pengirim adalah admin/owner
    member = await bot.get_chat_member(msg.chat.id, msg.from_user.id)
    if member.status not in ("administrator", "creator"):
        return await msg.reply("Hanya admin yang bisa export data\\.", parse_mode="MarkdownV2")

    now = now_wib()
    args = (msg.text or "").split()
    try:
        if len(args) > 1:
            year, month = map(int, args[1].split("-"))
        else:
            year, month = now.year, now.month
    except Exception:
        return await msg.reply("Format: `/export` atau `/export 2026\\-04`", parse_mode="MarkdownV2")

    rows = get_monthly_stats(msg.chat.id, year, month)
    if not rows:
        return await msg.reply(f"Belum ada data untuk {year}\\-{month:02d}\\.", parse_mode="MarkdownV2")

    # Build Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d}"

    header_fill = PatternFill("solid", fgColor="4F81BD")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["No", "Nama", "Username", "Tanggal", "Pesan"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i, r in enumerate(rows, 1):
        ws.append([
            i,
            r["full_name"] or "-",
            f"@{r['username']}" if r["username"] else "-",
            r["date"],
            r["total"],
        ])

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 8

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    month_esc = esc(f"{year}-{month:02d}")
    chat_name = esc(msg.chat.title or str(msg.chat.id))

    # Kirim ke DM admin
    try:
        await bot.send_document(
            msg.from_user.id,
            BufferedInputFile(buf.read(), f"statbot_{year}-{month:02d}.xlsx"),
            caption=f"📊 Export statistik *{chat_name}* bulan *{month_esc}*",
            parse_mode="MarkdownV2"
        )
        await msg.reply("File dikirim ke DM kamu\\.", parse_mode="MarkdownV2")
    except Exception:
        await msg.reply("Gagal kirim DM\\. Pastikan kamu sudah /start bot di private chat\\.", parse_mode="MarkdownV2")

# ── Record every message (registered last so commands take priority) ──
@dp.message(F.chat.type.in_({"group", "supergroup"}))
async def on_message(msg: Message):
    if not msg.from_user or msg.from_user.is_bot:
        return
    if msg.text and msg.text.startswith("/"):
        return
    dt = now_wib()
    record_message(
        chat_id=msg.chat.id,
        user_id=msg.from_user.id,
        username=msg.from_user.username or "",
        full_name=msg.from_user.full_name or "",
        date=dt.strftime("%Y-%m-%d"),
        hour=dt.hour,
    )

# ── Main ──────────────────────────────────────────────────────────────
OWNER_ID = 568033927

async def auto_backup():
    while True:
        await asyncio.sleep(6 * 3600)  # setiap 6 jam
        try:
            db_path = get_db_path()
            if os.path.exists(db_path):
                with open(db_path, "rb") as f:
                    data = f.read()
                ts = now_wib().strftime("%Y-%m-%d_%H%M")
                await bot.send_document(
                    OWNER_ID,
                    BufferedInputFile(data, f"statbot_{ts}.json"),
                    caption=f"🗄 Auto DB backup — {ts}"
                )
        except Exception:
            pass

async def main():
    init_db()
    await bot.set_my_commands([
        {"command": "start", "description": "Mulai bot"},
        {"command": "help", "description": "Daftar fitur"},
        {"command": "top", "description": "Top member paling aktif (30 hari)"},
        {"command": "stat", "description": "Statistik pesan kamu"},
        {"command": "grupstat", "description": "Statistik grup"},
        {"command": "grafik", "description": "Grafik aktivitas 7 hari"},
        {"command": "jam", "description": "Grafik aktivitas per jam"},
        {"command": "statadmin", "description": "Statistik admin bulan ini"},
        {"command": "export", "description": "Export data ke Excel (admin)"},
    ])
    print("StatBot jalan...")
    asyncio.create_task(auto_backup())
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
